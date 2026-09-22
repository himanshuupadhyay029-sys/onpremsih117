"""excel.py — Deterministic Excel Spreadsheet (.xlsx) builder for KAVACH.

The assigned local LLM (reasoning role) produces a clean, structured JSON description
of the spreadsheet (sheets, columns, rows, formulas, number formats), and openpyxl
deterministically compiles that JSON into a styled, professional .xlsx workbook.
"""

from datetime import datetime, timezone
import json
from pathlib import Path
import re
from typing import Any, Dict, List, Optional, Union

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False
    openpyxl = None
    Alignment = None
    Border = None
    Font = None
    PatternFill = None
    Side = None
    get_column_letter = None

from backend import config
from backend.audit.logbook import log_event
from backend.engine import active_llm as ollama, registry
from backend.terminal_logger import log_tool


def _log_terminal(msg: str) -> None:
    log_tool("excel", "XLSX", msg)


DRAFT_SPREADSHEET_PROMPT_TEMPLATE = """You are an expert industrial operations data analyst and spreadsheet engineer for KAVACH.
Based on the topic, calculations, or operational data below, design a clean, structured JSON blueprint for an Excel spreadsheet workbook (.xlsx).

Topic / Data Request:
{topic_or_content}

{sources_block}

CRITICAL RULES:
1. Return ONLY a valid JSON object matching this exact structure:
{{
  "title": "Short descriptive spreadsheet title",
  "filename_slug": "short_snake_case_name",
  "description": "Brief description of the data and purpose",
  "sheets": [
    {{
      "name": "Summary Sheet Name",
      "headers": ["Column 1", "Column 2", "Column 3", "Column 4"],
      "rows": [
        ["Sample Label 1", 120.50, 45, "=B2*C2"],
        ["Sample Label 2", 210.00, 30, "=B3*C3"]
      ],
      "summary_row": ["Total / Average", "=SUM(B2:B3)", "=AVERAGE(C2:C3)", "=SUM(D2:D3)"],
      "column_types": ["text", "currency", "number", "currency"]
    }}
  ]
}}

2. FORMULA USAGE: Use valid Excel uppercase formulas (e.g. =SUM(B2:B5), =AVERAGE(C2:C5), =B2*C2, =IF(D2>100, "HIGH", "NORMAL")).
3. NUMERIC ACCURACY: Use real numbers for values (not strings with commas), so Excel calculates properly.
4. COLUMN TYPES: Valid types are "text", "number", "currency", "percentage", "date".
5. Return ONLY the raw JSON object. No Markdown code fences, no conversational prose.
"""


def _sanitize_slug(name: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", name or "spreadsheet").strip().lower()
    return re.sub(r"[-\s]+", "_", cleaned)[:40] or "spreadsheet"


def _clean_json_response(raw: str) -> Dict[str, Any]:
    text = (raw or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    text = text.strip()

    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    match = re.search(r"(\{[\s\S]*\})", text)
    if match:
        try:
            return json.loads(match.group(1))
        except json.JSONDecodeError:
            pass

    return {
        "title": "Industrial Operations Spreadsheet",
        "filename_slug": "operations_data",
        "description": "Generated spreadsheet data",
        "sheets": [
            {
                "name": "Data",
                "headers": ["Item / Metric", "Value", "Notes"],
                "rows": [["Generated Entry", 100, "Extracted from task prompt"]],
                "summary_row": ["Total", "=SUM(B2:B2)", ""],
                "column_types": ["text", "number", "text"],
            }
        ],
    }


def draft_spreadsheet(
    topic_or_content: str,
    sources: Optional[List[Union[str, dict]]] = None,
    model: Optional[str] = None,
) -> Dict[str, Any]:
    """Uses the assigned LLM to produce a structured JSON blueprint for the spreadsheet."""
    resolved_model = model or registry.get_model("reasoning")
    _log_terminal(f"Drafting spreadsheet structure with model '{resolved_model}'...")

    sources_block = ""
    if sources:
        names = [s.get("filename", str(s)) if isinstance(s, dict) else str(s) for s in sources]
        sources_block = f"Available Sources:\n" + "\n".join(f"- {name}" for name in names)

    prompt = DRAFT_SPREADSHEET_PROMPT_TEMPLATE.format(
        topic_or_content=topic_or_content,
        sources_block=sources_block,
    )

    raw_response = ollama.generate(resolved_model, prompt)
    structured = _clean_json_response(raw_response)

    if not structured.get("sheets") or not isinstance(structured["sheets"], list):
        structured["sheets"] = [
            {
                "name": "Summary",
                "headers": ["Description", "Quantity", "Unit Rate", "Total"],
                "rows": [["Sample Item", 10, 250, "=B2*C2"]],
                "summary_row": ["Total", "=SUM(B2:B2)", "", "=SUM(D2:D2)"],
                "column_types": ["text", "number", "currency", "currency"],
            }
        ]

    return structured


def render_xlsx(spreadsheet_data: Dict[str, Any]) -> Path:
    """Deterministically compiles structured spreadsheet JSON into a professional styled .xlsx file."""
    config.OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)

    title = spreadsheet_data.get("title", "Spreadsheet")
    slug = _sanitize_slug(spreadsheet_data.get("filename_slug") or title)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    file_path = config.OUTPUTS_DIR / f"{slug}_{timestamp}.xlsx"

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    # Styling Palettes (Dark Navy Executive Industrial Theme)
    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    zebra_fill_even = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    zebra_fill_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    data_font = Font(name="Calibri", size=10, color="1E293B")
    data_alignment_left = Alignment(horizontal="left", vertical="center")
    data_alignment_right = Alignment(horizontal="right", vertical="center")
    data_alignment_center = Alignment(horizontal="center", vertical="center")

    summary_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    summary_font = Font(name="Calibri", size=11, bold=True, color="0F172A")

    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )
    summary_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="0F172A"),
        bottom=Side(style="double", color="0F172A"),
    )

    sheets_data = spreadsheet_data.get("sheets", [])
    if not sheets_data:
        sheets_data = [{"name": "Data", "headers": ["Item", "Value"], "rows": [["Sample", 1]]}]

    for idx, s_data in enumerate(sheets_data):
        sheet_name = re.sub(r"[\\/*?:\[\]]", "", s_data.get("name", f"Sheet{idx+1}"))[:31] or f"Sheet{idx+1}"
        ws = wb.create_sheet(title=sheet_name)

        headers = s_data.get("headers", [])
        rows = s_data.get("rows", [])
        summary_row = s_data.get("summary_row")
        col_types = s_data.get("column_types", ["text"] * len(headers))

        # 1. Header Row
        ws.row_dimensions[1].height = 28
        for col_num, header_text in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_num, value=str(header_text))
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = thin_border

        # 2. Data Rows
        current_row = 2
        for r_idx, row_vals in enumerate(rows):
            ws.row_dimensions[current_row].height = 20
            row_fill = zebra_fill_even if r_idx % 2 == 0 else zebra_fill_odd

            for col_num in range(1, len(headers) + 1):
                raw_val = row_vals[col_num - 1] if col_num - 1 < len(row_vals) else ""
                col_type = col_types[col_num - 1].lower() if col_num - 1 < len(col_types) else "text"
                cell = ws.cell(row=current_row, column=col_num)

                # Check if it is a formula
                if isinstance(raw_val, str) and raw_val.startswith("="):
                    cell.value = raw_val
                    cell.alignment = data_alignment_right
                    if "currency" in col_type:
                        cell.number_format = '"₹"#,##0.00;("₹"#,##0.00);"-"'
                    elif "percent" in col_type:
                        cell.number_format = "0.0%"
                    else:
                        cell.number_format = "#,##0.00"
                else:
                    # Parse numeric value if possible
                    parsed_num = None
                    if isinstance(raw_val, (int, float)):
                        parsed_num = raw_val
                    elif isinstance(raw_val, str) and re.match(r"^-?\d+(?:\.\d+)?$", raw_val.strip().replace(",", "")):
                        try:
                            clean_str = raw_val.strip().replace(",", "")
                            parsed_num = float(clean_str) if "." in clean_str else int(clean_str)
                        except Exception:
                            parsed_num = None

                    if parsed_num is not None:
                        cell.value = parsed_num
                        cell.alignment = data_alignment_right
                        if "currency" in col_type:
                            cell.number_format = '"₹"#,##0.00;("₹"#,##0.00);"-"'
                        elif "percent" in col_type:
                            cell.number_format = "0.0%"
                        elif isinstance(parsed_num, float):
                            cell.number_format = "#,##0.00"
                        else:
                            cell.number_format = "#,##0"
                    else:
                        cell.value = str(raw_val)
                        cell.alignment = data_alignment_left if col_type == "text" else data_alignment_center

                cell.fill = row_fill
                cell.font = data_font
                cell.border = thin_border

            current_row += 1

        # 3. Optional Summary Row
        if summary_row and isinstance(summary_row, list):
            ws.row_dimensions[current_row].height = 24
            for col_num in range(1, len(headers) + 1):
                sum_val = summary_row[col_num - 1] if col_num - 1 < len(summary_row) else ""
                col_type = col_types[col_num - 1].lower() if col_num - 1 < len(col_types) else "text"
                cell = ws.cell(row=current_row, column=col_num)

                if isinstance(sum_val, str) and sum_val.startswith("="):
                    cell.value = sum_val
                    cell.alignment = data_alignment_right
                    if "currency" in col_type:
                        cell.number_format = '"₹"#,##0.00;("₹"#,##0.00);"-"'
                    elif "percent" in col_type:
                        cell.number_format = "0.0%"
                    else:
                        cell.number_format = "#,##0.00"
                else:
                    cell.value = str(sum_val)
                    cell.alignment = data_alignment_left if col_num == 1 else data_alignment_right

                cell.fill = summary_fill
                cell.font = summary_font
                cell.border = summary_border

        # 4. Auto-fit column widths
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                if not val_str.startswith("="):
                    max_len = max(max_len, len(val_str))
                else:
                    max_len = max(max_len, 12)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Remove default empty sheet created by openpyxl
    if default_sheet in wb.worksheets and len(wb.worksheets) > 1:
        wb.remove(default_sheet)

    wb.save(str(file_path))
    _log_terminal(f"Successfully compiled Excel spreadsheet -> {file_path.name}")
    return file_path


def write_spreadsheet(
    prompt: str,
    sources: Optional[List[Union[str, dict]]] = None,
    task_id: Optional[str] = None,
    model: Optional[str] = None,
) -> Dict[str, Any]:
    """High-level end-to-end tool to draft, format, and save an .xlsx spreadsheet."""
    structured = draft_spreadsheet(prompt, sources=sources, model=model)
    file_path = render_xlsx(structured)
    filename = file_path.name

    total_rows = sum(len(s.get("rows", [])) for s in structured.get("sheets", []))

    log_event(
        task_id=task_id,
        event_type="write",
        actor="excel_tool",
        summary=f"Rendered Excel spreadsheet '{structured.get('title')}' ({len(structured.get('sheets', []))} sheets, {total_rows} rows) -> {filename}",
        metadata={
            "filename": filename,
            "file_path": str(file_path),
            "title": structured.get("title"),
            "sheets_count": len(structured.get("sheets", [])),
            "total_rows": total_rows,
        },
        external_calls=0,
    )

    return {
        "success": True,
        "title": structured.get("title", "Spreadsheet"),
        "filename": filename,
        "file_path": str(file_path),
        "download_url": f"/download/{filename}",
        "sheets": [s.get("name") for s in structured.get("sheets", [])],
        "total_rows": total_rows,
        "description": structured.get("description", ""),
    }
