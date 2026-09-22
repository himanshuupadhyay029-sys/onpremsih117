"""test_office_deliverables.py — Verifies Excel (.xlsx) and PowerPoint (.pptx) tools.

Tests:
1. Excel compilation with styles, number formats, and formulas via openpyxl.
2. PowerPoint compilation with 16:9 widescreen, custom themes, multiple layouts via python-pptx.
3. Router intent classification for spreadsheet and presentation prompts.
"""

import os
import sys
from pathlib import Path

# Add project root to sys.path
PROJECT_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(PROJECT_ROOT))

from backend.brain.router import route
from backend.tools.excel import render_xlsx
from backend.tools.ppt import render_pptx
import openpyxl
from pptx import Presentation


def test_excel_render():
    print("--- [TEST 1] Testing Excel .xlsx Rendering ---")
    mock_data = {
        "title": "Generator Monthly Fuel & Cost Analysis",
        "filename_slug": "test_generator_fuel_log",
        "sheets": [
            {
                "name": "Fuel Log",
                "headers": ["Date", "Generator ID", "Fuel Consumed (L)", "Rate (INR)", "Total Cost (INR)", "Status"],
                "rows": [
                    ["2026-09-01", "GEN-01", 120.5, 95.0, "=C2*D2", "Normal"],
                    ["2026-09-02", "GEN-01", 145.0, 95.0, "=C3*D3", "Heavy Load"],
                    ["2026-09-03", "GEN-02", 90.0, 95.0, "=C4*D4", "Normal"],
                ],
                "summary_row": ["Total / Average", "", "=SUM(C2:C4)", "=AVERAGE(D2:D4)", "=SUM(E2:E4)", ""],
                "column_types": ["date", "text", "number", "currency", "currency", "text"],
            }
        ],
    }

    file_path = render_xlsx(mock_data)
    assert file_path.exists(), f"File {file_path} was not created!"
    assert file_path.suffix == ".xlsx", f"Expected .xlsx file, got {file_path.suffix}"

    # Load back with openpyxl to verify integrity
    wb = openpyxl.load_workbook(str(file_path))
    assert "Fuel Log" in wb.sheetnames, f"Sheet 'Fuel Log' missing from {wb.sheetnames}"
    ws = wb["Fuel Log"]
    assert ws.cell(row=1, column=1).value == "Date"
    assert ws.cell(row=5, column=3).value == "=SUM(C2:C4)"
    print(f"✅ Excel rendering passed: {file_path.name} ({os.path.getsize(file_path)} bytes)")


def test_ppt_render():
    print("\n--- [TEST 2] Testing PowerPoint .pptx Rendering ---")
    mock_ppt = {
        "title": "Emergency Shutdown SOP Executive Briefing",
        "subtitle": "KAVACH Sovereign Operations Protocol",
        "filename_slug": "test_emergency_shutdown_briefing",
        "slides": [
            {
                "layout": "title",
                "title": "Emergency Shutdown SOP Executive Briefing",
                "subtitle": "Industrial Safety Directives | Zero Cloud Air-Gap Compliance",
            },
            {
                "layout": "content",
                "title": "Executive Summary & Objectives",
                "summary": "Mandatory procedure for rapid isolation of fuel lines and electrical grids.",
                "bullets": [
                    "Immediate trip threshold: Turbine vibration > 4.5 mm/s or Boiler pressure > 15.2 bar.",
                    "Solenoid cut-off response target: < 200 ms under autonomous controller action.",
                    "Zero external network exposure ensures immunity from cyber-physical disruption.",
                ],
                "notes": "Speaker notes: Emphasize that all control loops execute on local hardware.",
            },
            {
                "layout": "two_column",
                "title": "Current State vs KAVACH Directive",
                "left_heading": "Legacy Protocol",
                "left_bullets": [
                    "Manual valve actuation taking 4-8 minutes.",
                    "Paper logbook recording with delayed shift handover.",
                ],
                "right_heading": "KAVACH Sovereign SOP",
                "right_bullets": [
                    "Autonomous sensor telemetry tripping in milliseconds.",
                    "Cryptographic append-only local audit logbook.",
                ],
                "notes": "Compare response times and reliability.",
            },
            {
                "layout": "metrics_summary",
                "title": "Operational KPIs & Action Plan",
                "metrics": [
                    {"label": "Trip Latency", "value": "180 ms"},
                    {"label": "Compliance", "value": "100%"},
                    {"label": "Air-Gap Level", "value": "Tier 4"},
                ],
                "action_items": [
                    "Inspect mechanical interlocks every Monday shift.",
                    "Verify local emergency siren batteries quarterly.",
                ],
            },
        ],
    }

    file_path = render_pptx(mock_ppt)
    assert file_path.exists(), f"File {file_path} was not created!"
    assert file_path.suffix == ".pptx", f"Expected .pptx file, got {file_path.suffix}"

    # Load back with python-pptx to verify integrity
    prs = Presentation(str(file_path))
    assert len(prs.slides) == 4, f"Expected 4 slides, got {len(prs.slides)}"
    assert abs(prs.slide_width.inches - 13.333) < 0.01, "Expected 16:9 widescreen width"
    print(f"✅ PowerPoint rendering passed: {file_path.name} (4 slides, {os.path.getsize(file_path)} bytes)")


def test_router_classification():
    print("\n--- [TEST 3] Testing Router Classification for Excel & PPT ---")
    
    dec_excel = route("Create a monthly generator fuel consumption log spreadsheet with sum formulas")
    assert dec_excel.task_type == "excel", f"Expected 'excel', got {dec_excel.task_type}"
    print(f"✅ Router: 'Create spreadsheet...' -> '{dec_excel.task_type}' (model role: '{dec_excel.model_role}')")

    dec_ppt = route("Generate a 4-slide executive presentation briefing on the plant emergency shutdown procedure")
    assert dec_ppt.task_type == "ppt", f"Expected 'ppt', got {dec_ppt.task_type}"
    print(f"✅ Router: 'Generate 4-slide presentation...' -> '{dec_ppt.task_type}' (model role: '{dec_ppt.model_role}')")


if __name__ == "__main__":
    test_excel_render()
    test_ppt_render()
    test_router_classification()
    print("\n🎉 ALL OFFICE DELIVERABLE TESTS PASSED!")
